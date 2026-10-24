import sys
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import MagicMock, patch

from docx import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pydantic import ValidationError

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.ingestion.loaders.docx_loader import DocxLoader
from src.ingestion.loaders.pdf_loader import PDFLoader
from src.schemas.blocks import BlockType, ParsedBlock, TableData


class TestParsedBlockSchema(unittest.TestCase):
    def test_table_data_preserves_header_rows_caption_and_ragged_cells(self):
        table = TableData.from_matrix(
            [["Mô hình", "F1"], ["Random Forest", "0.9752"], ["XGBoost"]],
            caption="Bảng 5: Kết quả",
            table_id="table-5",
        )

        self.assertEqual(table.table_id, "table-5")
        self.assertEqual(table.caption, "Bảng 5: Kết quả")
        self.assertEqual([column.name for column in table.columns], ["Mô hình", "F1"])
        self.assertEqual(table.header, ["Mô hình", "F1"])
        self.assertEqual(table.rows, [["Random Forest", "0.9752"], ["XGBoost"]])

    def test_block_validation_requires_structural_table_and_heading_level(self):
        with self.assertRaises(ValidationError):
            ParsedBlock(block_index=0, block_type=BlockType.TABLE)
        with self.assertRaises(ValidationError):
            ParsedBlock(block_index=0, block_type=BlockType.HEADING, text="Chương 1")
        with self.assertRaises(ValidationError):
            ParsedBlock(
                block_index=0,
                block_type=BlockType.PARAGRAPH,
                text="Nội dung",
                heading_level=1,
            )

    def test_section_identifier_is_source_scoped_and_stable(self):
        path = ["Chương 1", "1.1 Phạm vi"]
        self.assertEqual(
            ParsedBlock.section_identifier("source-a", path),
            ParsedBlock.section_identifier("source-a", path),
        )
        self.assertNotEqual(
            ParsedBlock.section_identifier("source-a", path),
            ParsedBlock.section_identifier("source-b", path),
        )


class TestLoaderParsedBlocks(unittest.TestCase):
    def _load_docx_blocks(self, source, *, loader=None, indexable_only=False):
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.docx"
            source.save(path)
            return (loader or DocxLoader()).load_blocks(path, indexable_only=indexable_only)

    def test_pdf_loader_exposes_structured_table_and_real_page_number(self):
        class FakePage:
            chars = []

        table = TableData.from_matrix(
            [["Epoch", "Accuracy"], ["10", "0.90"]],
            caption="Bảng 1: Kết quả",
        )
        source_blocks = [
            {"type": "heading", "y": 10, "content": "CHƯƠNG 1", "level": 1},
            {
                "type": "table",
                "y": 20,
                "table": table,
                "bbox": (10, 20, 100, 120),
            },
        ]
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.pdf"
            path.touch()
            pdf_context = MagicMock()
            pdf_context.__enter__.return_value.pages = [FakePage()]
            with (
                patch("src.ingestion.loaders.pdf_loader.pdfplumber.open", return_value=pdf_context),
                patch.object(PDFLoader, "_page_to_blocks", return_value=source_blocks),
            ):
                blocks = PDFLoader(ocr_enabled=False).load_blocks(path)

        self.assertEqual([block.block_type for block in blocks], [BlockType.HEADING, BlockType.TABLE])
        self.assertEqual(blocks[1].page_number, 1)
        self.assertEqual(blocks[1].bbox.x0, 10)
        self.assertEqual(blocks[1].table.header, ["Epoch", "Accuracy"])
        self.assertEqual(blocks[1].table.rows, [["10", "0.90"]])

    def test_docx_loader_exposes_table_without_invented_page_number(self):
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.docx"
            source = DocxDocument()
            source.add_heading("Kết quả", level=1)
            table = source.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Mô hình"
            table.rows[0].cells[1].text = "F1"
            row = table.add_row().cells
            row[0].text = "Random Forest"
            row[1].text = "0.9752"
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            table.rows[0]._tr.get_or_add_trPr().append(header)
            source.save(path)

            loader = DocxLoader()
            blocks = loader.load_blocks(path)
            legacy_documents = loader.load(path)

        table_block = next(block for block in blocks if block.block_type == BlockType.TABLE)
        self.assertIsNone(table_block.page_number)
        self.assertEqual(table_block.table.header, ["Mô hình", "F1"])
        self.assertEqual(table_block.table.header_rows, [["Mô hình", "F1"]])
        self.assertEqual(table_block.table.rows, [["Random Forest", "0.9752"]])
        self.assertIn("parsed_block_ids", legacy_documents[-1].metadata)
        self.assertIn("| Mô hình | F1 |", legacy_documents[-1].text)

    def test_docx_heading_detection_prioritizes_styles_and_outline_levels(self):
        source = DocxDocument()
        source.add_paragraph("Styled chapter", style="Heading 1")
        source.add_paragraph("Styled section", style="Heading 2")
        outline = source.add_paragraph("Outline subsection")
        outline_properties = outline._p.get_or_add_pPr()
        outline_level = OxmlElement("w:outlineLvl")
        outline_level.set(qn("w:val"), "2")
        outline_properties.append(outline_level)
        styled_outline = source.add_paragraph("Style takes precedence", style="Heading 2")
        styled_outline_properties = styled_outline._p.get_or_add_pPr()
        conflicting_outline_level = OxmlElement("w:outlineLvl")
        conflicting_outline_level.set(qn("w:val"), "0")
        styled_outline_properties.append(conflicting_outline_level)

        headings = [
            block
            for block in self._load_docx_blocks(source)
            if block.block_type == BlockType.HEADING
        ]

        self.assertEqual(
            [(block.text, block.heading_level) for block in headings],
            [
                ("Styled chapter", 1),
                ("Styled section", 2),
                ("Outline subsection", 3),
                ("Style takes precedence", 2),
            ],
        )

    def test_docx_heading_detection_uses_numbering_and_safe_direct_formatting_fallback(self):
        source = DocxDocument()
        numbered_heading = source.add_paragraph()
        numbered_run = numbered_heading.add_run("2. Scope")
        numbered_run.bold = True
        numbered_run.font.size = Pt(16)
        source.add_paragraph("2.1 Dataset")
        source.add_paragraph("2.1.1 Cleaning")

        direct_heading = source.add_paragraph()
        run = direct_heading.add_run("Unnumbered bold heading")
        run.bold = True
        run.font.size = Pt(16)

        numbered_list = source.add_paragraph("2.1 This is a list item", style="List Paragraph")
        numbered_list_properties = numbered_list._p.get_or_add_pPr()
        numbered_list_properties.append(OxmlElement("w:numPr"))
        source.add_paragraph("1. A plain numbered list item")

        uppercase = source.add_paragraph()
        uppercase_run = uppercase.add_run("BACKGROUND AND MOTIVATION")
        uppercase_run.bold = True
        uppercase_run.font.size = Pt(20)

        sentence = source.add_paragraph()
        sentence_run = sentence.add_run("This is an ordinary sentence.")
        sentence_run.bold = True
        sentence_run.font.size = Pt(20)
        source.add_paragraph("Chapter 3 explains the method.")

        blocks = self._load_docx_blocks(source)
        headings = [
            (block.text, block.heading_level)
            for block in blocks
            if block.block_type == BlockType.HEADING
        ]
        non_headings = {block.text for block in blocks if block.block_type == BlockType.PARAGRAPH}

        self.assertEqual(
            headings,
            [
                ("2. Scope", 1),
                ("2.1 Dataset", 2),
                ("2.1.1 Cleaning", 3),
                ("Unnumbered bold heading", 1),
            ],
        )
        self.assertIn("2.1 This is a list item", non_headings)
        self.assertIn("1. A plain numbered list item", non_headings)
        self.assertIn("BACKGROUND AND MOTIVATION", non_headings)
        self.assertIn("This is an ordinary sentence.", non_headings)
        self.assertIn("Chapter 3 explains the method.", non_headings)

    def test_docx_regions_exclude_toc_entries_and_reset_breadcrumbs(self):
        source = DocxDocument()
        source.add_paragraph("University")
        source.add_paragraph("Project title", style="Title")
        source.add_paragraph("MỤC LỤC", style="Heading 1")

        toc_style = source.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
        source.add_paragraph("1. Introduction", style=toc_style)
        field_paragraph = source.add_paragraph()
        field_paragraph.add_run("2. Results")
        toc_field = OxmlElement("w:fldSimple")
        toc_field.set(qn("w:instr"), 'TOC \\o "1-3"')
        field_run = OxmlElement("w:r")
        field_text = OxmlElement("w:t")
        field_text.text = "2. Results"
        field_run.append(field_text)
        toc_field.append(field_run)
        field_paragraph._p.append(toc_field)
        source.add_paragraph("TÀI LIỆU THAM KHẢO\t44", style=toc_style)

        source.add_paragraph("DANH MỤC HÌNH", style="Heading 1")
        source.add_paragraph("Hình 1. Sơ đồ", style=toc_style)
        source.add_paragraph("DANH MỤC BẢNG", style="Heading 1")
        source.add_paragraph("Bảng 1. Kết quả", style=toc_style)
        source.add_paragraph("MỞ ĐẦU", style="Heading 1")
        source.add_paragraph("Chapter 1: Introduction", style="Heading 1")
        source.add_paragraph("1.1 Background", style="Heading 2")
        source.add_paragraph("Chapter 2: Experiments", style="Heading 1")
        source.add_paragraph("Fresh chapter content.")
        source.add_paragraph("TÀI LIỆU THAM KHẢO", style="Heading 1")

        blocks = self._load_docx_blocks(source)
        by_text = {block.text: block for block in blocks if block.text}

        self.assertEqual(by_text["University"].region.value, "cover")
        self.assertEqual(by_text["MỤC LỤC"].region.value, "table_of_contents")
        self.assertEqual(by_text["1. Introduction"].region.value, "table_of_contents")
        self.assertEqual(by_text["2. Results"].region.value, "table_of_contents")
        self.assertEqual(by_text["TÀI LIỆU THAM KHẢO\t44"].region.value, "table_of_contents")
        self.assertEqual(by_text["DANH MỤC HÌNH"].region.value, "list_of_figures")
        self.assertEqual(by_text["DANH MỤC BẢNG"].region.value, "list_of_tables")
        self.assertEqual(by_text["MỞ ĐẦU"].region.value, "body")
        self.assertEqual(by_text["Chapter 1: Introduction"].region.value, "body")
        self.assertEqual(by_text["TÀI LIỆU THAM KHẢO"].region.value, "references")
        self.assertEqual(by_text["Fresh chapter content."].section_path, ["Chapter 2: Experiments"])
        self.assertNotIn("Chapter 1: Introduction", by_text["Fresh chapter content."].section_path)

        indexable = self._load_docx_blocks(
            source,
            loader=DocxLoader(index_regions={"body"}),
            indexable_only=True,
        )
        self.assertTrue(indexable)
        self.assertTrue(all(block.region.value == "body" for block in indexable))
        self.assertNotIn("TÀI LIỆU THAM KHẢO", {block.text for block in indexable})

    def test_docx_tables_preserve_order_merges_headers_and_captions(self):
        source = DocxDocument()
        source.add_paragraph("Before table.")
        source.add_table(rows=1, cols=2)
        punctuation_table = source.add_table(rows=1, cols=2)
        punctuation_table.cell(0, 0).text = "—"
        punctuation_table.cell(0, 1).text = "..."

        caption_text = "Bảng 1: Kết quả thực nghiệm"
        source.add_paragraph(caption_text)
        table = source.add_table(rows=4, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Phân loại"
        table.cell(0, 2).text = "Kết quả"
        table.cell(1, 0).text = "Nhóm"
        table.cell(1, 1).text = "Mô hình"
        table.cell(1, 2).text = "F1"
        table.cell(2, 0).text = "Mô hình"
        table.cell(2, 1).text = "Random Forest"
        table.cell(2, 2).text = "0.95"
        table.cell(3, 1).text = "XGBoost"
        table.cell(3, 2).text = "0.97"
        table.cell(2, 0).merge(table.cell(3, 0))

        for row in table.rows[:2]:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            row._tr.get_or_add_trPr().append(header)

        source.add_paragraph("After table.")
        following_caption_table = source.add_table(rows=1, cols=2)
        following_caption_table.cell(0, 0).text = "DenseNet201"
        following_caption_table.cell(0, 1).text = "0.94"
        following_caption_text = "Bảng 2: Kết quả DenseNet201"
        source.add_paragraph(following_caption_text)

        uncaptioned = source.add_table(rows=1, cols=2)
        uncaptioned.cell(0, 0).text = "Left value"
        uncaptioned.cell(0, 1).text = "Right value"

        blocks = self._load_docx_blocks(source)
        self.assertEqual(
            [block.block_type for block in blocks],
            [
                BlockType.PARAGRAPH,
                BlockType.CAPTION,
                BlockType.TABLE,
                BlockType.PARAGRAPH,
                BlockType.TABLE,
                BlockType.CAPTION,
                BlockType.TABLE,
            ],
        )
        first_table = blocks[2]
        self.assertEqual(first_table.table.caption, caption_text)
        self.assertEqual(first_table.table.header_rows, [["Phân loại", "", "Kết quả"], ["Nhóm", "Mô hình", "F1"]])
        self.assertEqual(first_table.table.header, ["Nhóm", "Mô hình", "F1"])
        self.assertEqual(first_table.table.rows, [["Mô hình", "Random Forest", "0.95"], ["", "XGBoost", "0.97"]])
        self.assertEqual(first_table.table.cells[0][0].column_span, 2)
        self.assertEqual(first_table.table.cells[2][0].row_span, 2)
        self.assertEqual(first_table.table.cells[3][0].vertical_merge, "continue")
        self.assertEqual(
            first_table.metadata["caption_block_id"],
            blocks[1].block_id,
        )
        following_caption_table_block = blocks[4]
        following_caption_block = blocks[5]
        self.assertEqual(following_caption_table_block.table.caption, following_caption_text)
        self.assertEqual(
            following_caption_table_block.metadata["caption_block_id"],
            following_caption_block.block_id,
        )
        self.assertEqual(
            following_caption_block.metadata["table_block_id"],
            following_caption_table_block.block_id,
        )
        self.assertTrue(blocks[6].metadata["caption_missing"])
        self.assertIsNone(blocks[6].table.caption)
        self.assertEqual(blocks[6].table.header, [])
        self.assertEqual(blocks[6].table.rows, [["Left value", "Right value"]])

        legacy_documents = DocxLoader().load(
            self._save_docx_and_return_path(source)
        )
        rendered = "\n".join(document.text for document in legacy_documents)
        self.assertEqual(rendered.count(caption_text), 1)
        self.assertEqual(rendered.count(following_caption_text), 1)
        self.assertNotIn("Bảng (thuộc:", rendered)

    def _save_docx_and_return_path(self, source):
        temp_dir = TemporaryDirectory()
        self.addCleanup(temp_dir.cleanup)
        path = Path(temp_dir.name) / "sample.docx"
        source.save(path)
        return path


if __name__ == "__main__":
    unittest.main()
