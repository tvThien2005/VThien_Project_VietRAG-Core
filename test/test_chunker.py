import sys
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from docx import Document as DocxDocument

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.ingestion.chunkers.recursive_chunker import RecursiveCharacterChunker
from src.ingestion.loaders.docx_loader import DocxLoader
from src.ingestion.loaders.pdf_loader import (
    PDFLoader,
    _classify_line,
    _extract_heading_level,
)
from src.schemas.document import Document
from src.utils.text import clean_text, normalize_vietnamese_unicode


class TestVietnameseChunker(unittest.TestCase):
    def setUp(self):
        self.chunker = RecursiveCharacterChunker(chunk_size=100, chunk_overlap=20)

    def test_unicode_normalization(self):
        decomposed = "trường đại học"  # NFD
        normalized = normalize_vietnamese_unicode(decomposed)
        self.assertEqual(len(normalized), len("trường đại học"))

    def test_clean_text(self):
        raw = "  Đoạn văn   tiếng   Việt.\n\n\nCó  nhiều   khoảng trắng.  "
        cleaned = clean_text(raw)
        self.assertNotIn("   ", cleaned)
        self.assertNotIn("\n\n\n", cleaned)

    def test_recursive_split(self):
        text = (
            "Hệ thống RAG cho phép tìm kiếm tài liệu chính xác. "
            "Sau khi tìm kiếm, mô hình ngôn ngữ lớn sẽ tổng hợp câu trả lời dựa trên ngữ cảnh được cung cấp. "
            "Tiếng Việt có các đặc thù ngữ pháp riêng cần lưu ý khi tách câu."
        )
        doc = Document(
            text=text,
            metadata={
                "file_name": "test.txt",
                "chapter": "Chương 1: Tổng quan",
                "section_title": "1.1 Phạm vi",
                "page_number": 3,
            },
        )
        chunks = self.chunker.split_document(doc)

        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertLessEqual(chunk.char_count, 120)
            self.assertEqual(chunk.metadata["file_name"], "test.txt")
            self.assertEqual(chunk.metadata["chapter"], "Chương 1: Tổng quan")
            self.assertEqual(chunk.metadata["section_title"], "1.1 Phạm vi")
            self.assertEqual(chunk.metadata["page_number"], 3)
            self.assertEqual(chunk.metadata["chunk_chapter"], "Chương 1: Tổng quan")
            self.assertEqual(chunk.metadata["chunk_section_title"], "1.1 Phạm vi")

    def test_pdf_loader_tracks_chapter_section_and_page(self):
        class FakePage:
            chars = []

        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "hierarchy.pdf"
            path.touch()
            pdf_context = unittest.mock.MagicMock()
            pdf_context.__enter__.return_value.pages = [FakePage()]
            blocks = [
                {"type": "heading", "y": 10, "content": "CHƯƠNG 1: TỔNG QUAN", "level": 1},
                {"type": "heading", "y": 20, "content": "1.1 Phạm vi", "level": 2},
                {"type": "text", "y": 30, "content": "Nội dung mục."},
            ]
            with (
                patch("src.ingestion.loaders.pdf_loader.pdfplumber.open", return_value=pdf_context),
                patch.object(PDFLoader, "_page_to_blocks", return_value=blocks),
            ):
                documents = PDFLoader(ocr_enabled=False).load(path)

        self.assertEqual(len(documents), 2)
        metadata = documents[-1].metadata
        self.assertEqual(metadata["chapter"], "CHƯƠNG 1: TỔNG QUAN")
        self.assertEqual(metadata["section_title"], "1.1 Phạm vi")
        self.assertEqual(metadata["breadcrumb"], "CHƯƠNG 1: TỔNG QUAN > 1.1 Phạm vi")
        self.assertEqual(metadata["page_number"], 1)

    def test_pdf_heading_classifier_ignores_bold_enumeration(self):
        chars = [{"size": 12, "fontname": "Arial-Bold"}]
        with patch(
            "src.ingestion.loaders.pdf_loader.pdfplumber.utils.extract_text",
            return_value="Thứ nhất, về chỉ số Silhouette:",
        ):
            is_heading, _ = _classify_line(chars, median_size=12)

        self.assertFalse(is_heading)

    def test_pdf_heading_classifier_accepts_lettered_heading_and_uses_subsection_level(self):
        chars = [{"size": 12, "fontname": "Arial-Bold"}]
        with patch(
            "src.ingestion.loaders.pdf_loader.pdfplumber.utils.extract_text",
            return_value="F. Kết quả huấn luyện và so sánh 2 thuật toán",
        ):
            is_heading, text = _classify_line(chars, median_size=12)

        self.assertTrue(is_heading)
        self.assertEqual(_extract_heading_level(text), 2)
        self.assertEqual(_extract_heading_level("CHƯƠNG 4: Tập dữ liệu Fitness"), 1)

    def test_pdf_loader_keeps_chapter_for_lettered_heading(self):
        class FakePage:
            chars = []

        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "hierarchy.pdf"
            path.touch()
            pdf_context = unittest.mock.MagicMock()
            pdf_context.__enter__.return_value.pages = [FakePage()]
            blocks = [
                {
                    "type": "heading",
                    "y": 10,
                    "content": "CHƯƠNG 4: Tập dữ liệu Fitness",
                    "level": 1,
                },
                {
                    "type": "heading",
                    "y": 20,
                    "content": "F. Kết quả huấn luyện và so sánh 2 thuật toán",
                    "level": 2,
                },
                {
                    "type": "text",
                    "y": 30,
                    "content": "Thứ nhất, về chỉ số Silhouette: K-Means đạt 0.4814.",
                },
            ]
            with (
                patch("src.ingestion.loaders.pdf_loader.pdfplumber.open", return_value=pdf_context),
                patch.object(PDFLoader, "_page_to_blocks", return_value=blocks),
            ):
                documents = PDFLoader(ocr_enabled=False).load(path)

        metadata = documents[-1].metadata
        self.assertEqual(metadata["chapter"], "CHƯƠNG 4: Tập dữ liệu Fitness")
        self.assertEqual(
            metadata["section_title"],
            "F. Kết quả huấn luyện và so sánh 2 thuật toán",
        )
        self.assertEqual(
            metadata["breadcrumb"],
            "CHƯƠNG 4: Tập dữ liệu Fitness > F. Kết quả huấn luyện và so sánh 2 thuật toán",
        )
        self.assertIn("Thứ nhất, về chỉ số Silhouette", documents[-1].text)

    def test_table_chunks_keep_caption_header_and_complete_rows(self):
        rows = [f"Mã {index} | Giá trị đặc trưng {index}" for index in range(1, 9)]
        table = "\n\n".join(
            ["Bảng 1: Kết quả đo", "Tên chỉ số | Giá trị\n--- | ---\n" + "\n".join(rows)]
        )
        chunks = self.chunker.split_document(Document(text=table))

        table_chunks = [chunk for chunk in chunks if chunk.metadata["chunk_has_table"]]
        self.assertGreater(len(table_chunks), 1)
        for chunk in table_chunks:
            self.assertIn("Bảng 1: Kết quả đo", chunk.text)
            self.assertIn("Tên chỉ số | Giá trị", chunk.text)
            self.assertIn("--- | ---", chunk.text)
            self.assertLessEqual(chunk.text.count(" | "), 3)
        for row in rows:
            self.assertEqual(sum(row in chunk.text for chunk in table_chunks), 1)

    def test_table_chunks_keep_caption_after_table_with_data(self):
        rows = [f"| {index} | Giá trị {index} |" for index in range(1, 9)]
        table = "\n".join(
            ["| Epoch | Accuracy |", "| --- | --- |", *rows, "", "Bảng 4: Kết quả huấn luyện"]
        )
        chunks = self.chunker.split_documents([Document(text=table)])

        table_chunks = [chunk for chunk in chunks if chunk.metadata["chunk_has_table"]]
        self.assertGreater(len(table_chunks), 1)
        for chunk in table_chunks:
            self.assertIn("Bảng 4: Kết quả huấn luyện", chunk.text)
            self.assertIn("| Epoch | Accuracy |", chunk.text)
            self.assertIn("| --- | --- |", chunk.text)
        for row in rows:
            self.assertEqual(sum(row in chunk.text for chunk in table_chunks), 1)

    def test_docx_single_column_table_keeps_caption_with_data(self):
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "table.docx"
            source = DocxDocument()
            source.add_paragraph("Bảng 1: Kết quả đo")
            table = source.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Tên chỉ số"
            rows = [f"Giá trị {index}" for index in range(1, 9)]
            for value in rows:
                table.add_row().cells[0].text = value
            source.save(path)

            documents = DocxLoader().load(path)
            chunks = self.chunker.split_documents(documents)

        table_chunks = [chunk for chunk in chunks if chunk.metadata["chunk_has_table"]]
        self.assertGreater(len(table_chunks), 1)
        for chunk in table_chunks:
            self.assertIn("Bảng 1: Kết quả đo", chunk.text)
            self.assertIn("Tên chỉ số", chunk.text)
            self.assertIn("| --- |", chunk.text)
        for row in rows:
            self.assertEqual(sum(row in chunk.text for chunk in table_chunks), 1)

    def test_pdf_single_column_table_keeps_bold_caption_with_data(self):
        class FoundTable:
            bbox = (0, 20, 100, 120)

            def extract(self):
                return [["Tên chỉ số"], *[[f"Giá trị {index}"] for index in range(1, 9)]]

        class FakePage:
            chars = [{
                "x0": 1,
                "x1": 2,
                "top": 10,
                "bottom": 11,
                "size": 14,
                "fontname": "Arial-Bold",
            }]

            def extract_tables(self, settings):
                return [["Tên chỉ số"], ["Giá trị 1"]]

            def find_tables(self, settings):
                if settings["vertical_strategy"] == "lines":
                    return []
                return [FoundTable()]

        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "table.pdf"
            path.touch()
            pdf_context = unittest.mock.MagicMock()
            pdf_context.__enter__.return_value.pages = [FakePage()]
            with (
                patch("src.ingestion.loaders.pdf_loader.pdfplumber.open", return_value=pdf_context),
                patch(
                    "src.ingestion.loaders.pdf_loader.pdfplumber.utils.extract_text",
                    return_value="Bảng 1: Caption",
                ),
            ):
                documents = PDFLoader().load(path)

        chunks = self.chunker.split_documents(documents)
        table_chunks = [chunk for chunk in chunks if chunk.metadata["chunk_has_table"]]
        self.assertGreater(len(table_chunks), 1)
        for chunk in table_chunks:
            self.assertIn("Bảng 1: Caption", chunk.text)
            self.assertIn("Tên chỉ số", chunk.text)
            self.assertIn("| --- |", chunk.text)
        for index in range(1, 9):
            row = f"Giá trị {index}"
            self.assertEqual(sum(row in chunk.text for chunk in table_chunks), 1)

    def test_pdf_table_keeps_caption_after_data(self):
        class FoundTable:
            bbox = (0, 20, 100, 120)

            def extract(self):
                return [["Tên chỉ số"], *[[f"Giá trị {index}"] for index in range(1, 9)]]

        class FakePage:
            chars = [{
                "x0": 1,
                "x1": 2,
                "top": 130,
                "bottom": 131,
                "size": 14,
                "fontname": "Arial-Bold",
            }]

            def find_tables(self, settings):
                if settings["vertical_strategy"] == "lines":
                    return []
                return [FoundTable()]

        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "table.pdf"
            path.touch()
            pdf_context = unittest.mock.MagicMock()
            pdf_context.__enter__.return_value.pages = [FakePage()]
            with (
                patch("src.ingestion.loaders.pdf_loader.pdfplumber.open", return_value=pdf_context),
                patch(
                    "src.ingestion.loaders.pdf_loader.pdfplumber.utils.extract_text",
                    return_value="Bảng 1: Caption",
                ),
            ):
                documents = PDFLoader().load(path)

        chunks = self.chunker.split_documents(documents)
        table_chunks = [chunk for chunk in chunks if chunk.metadata["chunk_has_table"]]
        self.assertGreater(len(table_chunks), 1)
        for chunk in table_chunks:
            self.assertIn("Bảng 1: Caption", chunk.text)
            self.assertIn("Tên chỉ số", chunk.text)
            self.assertIn("| --- |", chunk.text)
        for index in range(1, 9):
            row = f"Giá trị {index}"
            self.assertEqual(sum(row in chunk.text for chunk in table_chunks), 1)

    def test_pdf_scanned_page_uses_ocr_fallback(self):
        class FakePage:
            chars = []

            def find_tables(self, settings):
                return []

        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "scanned.pdf"
            path.touch()
            pdf_context = unittest.mock.MagicMock()
            pdf_context.__enter__.return_value.pages = [FakePage()]
            loader = PDFLoader(ocr_enabled=True)
            with (
                patch("src.ingestion.loaders.pdf_loader.pdfplumber.open", return_value=pdf_context),
                patch.object(
                    loader,
                    "_ocr_page",
                    return_value="Tài liệu được nhận dạng từ ảnh.",
                ) as ocr_page,
            ):
                documents = loader.load(path)

        ocr_page.assert_called_once_with(path.resolve(), 0)
        self.assertEqual(len(documents), 1)
        self.assertIn("Tài liệu được nhận dạng từ ảnh.", documents[0].text)
        self.assertTrue(documents[0].metadata["ocr_used"])


if __name__ == "__main__":
    unittest.main()
